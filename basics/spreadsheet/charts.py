import openpyxl as xl
from openpyxl.chart import BarChart, Reference, Series

filenames= input("Enter the name of the file-{extension}")
def chart_generator(filenames):
    wb = xl.load_workbook(f"{filenames}.xlsx")
    sheet1=wb['Sheet1']

    chart = BarChart()
    values= Reference(sheet1, min_row= 2, max_row= sheet1.max_row, min_col=1, max_col=1)

    chart.add_data(values)
    sheet1.add_chart(chart, 'c2')

    wb.save(f"{filenames}.xlsx")

chart_generator(filenames)